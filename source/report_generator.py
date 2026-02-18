import json
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    Alignment,
    PatternFill,
    Border,
    Side,
    NamedStyle,
)

from openpyxl.utils import get_column_letter

from source.camera_info_grabber import BASE_DIR, CONFIG_DIR


reports_dir = BASE_DIR / "reports"
reports_dir.mkdir(exist_ok=True)


def generate_report(data: pd.DataFrame, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = datetime.now().strftime("%d.%m.%Y")

    f_bold = Font(bold=True)
    f_title = Font(size=16, bold=True)
    align_center = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )

    bg_gray = PatternFill("solid", fgColor="DDDDDD")
    bg_red = PatternFill("solid", fgColor="F4CCCC")

    thin_line = Side(style="thin")
    thick_line = Side(style="thick")
    thin_border = Border(
        left=thin_line, right=thin_line, top=thin_line, bottom=thin_line
    )
    thick_border = Border(
        left=thick_line, right=thick_line, top=thick_line, bottom=thick_line
    )

    cell_style = NamedStyle(name="cell_style")
    cell_style.font = Font(size=14)
    cell_style.border = thin_border
    wb.add_named_style(cell_style)

    ws.column_dimensions["A"].width = 36
    ws.row_dimensions[1].height = 60
    for col in range(2, 8):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 12

    ws.merge_cells("A1:G1")
    ws["A1"] = (
        "Відомості щодо стану функціонування системи відеоспостереження "
        'та відеоаналітики ("Безпечна Чернігівщина")'
    )
    ws["A1"].font = f_title
    ws["A1"].alignment = align_center
    ws["A1"].border = Border(
        left=thick_line, top=thick_line, bottom=thin_line
    )
    ws["G1"].border = Border(
        right=thick_line, top=thick_line, bottom=thin_line
    )
    for col in range(2, 7):
        ws.cell(row=1, column=col).border = Border(
            top=thick_line, bottom=thin_line
        )

    ws["A2"] = datetime.now().strftime("%d.%m.%Y")
    ws["A2"].font = Font(size=14, bold=True)

    ws.merge_cells("B2:D2")
    ws["B2"] = "Встановлено"
    ws["B2"].font = Font(size=14, bold=True)

    ws.merge_cells("E2:G2")
    ws["E2"] = "Не працює"
    ws["E2"].font = Font(size=14, bold=True)
    ws["E2"].fill = bg_red

    ws["A2"].border = Border(left=thick_line)
    ws["G2"].border = Border(right=thick_line)

    headers = ["всього", "номерні", "обличчя", "всього", "номерні", "обличчя"]
    ws["A3"] = "Населений пункт"
    ws["A3"].font = Font(size=14, bold=True)
    ws["A3"].border = Border(left=thick_line, bottom=thick_line)
    ws["G3"].border = Border(right=thick_line, bottom=thick_line)
    for i, value in enumerate(headers):
        ws.cell(row=3, column=2 + i, value=value).font = Font(
            size=14, bold=True
        )
        if i > 2:
            ws.cell(row=3, column=2 + i).fill = bg_red
    for col in range(2, 7):
        ws.cell(row=3, column=col).border = Border(bottom=thick_line)

    for row in ws.iter_rows(min_row=1, max_row=3, min_col=1, max_col=7):
        for cell in row:
            cell.alignment = align_center

    current_row = 4
    start_row = 4

    with open(CONFIG_DIR / "region.json", "r", encoding="utf-8") as f:
        districts_map = json.load(f)

    # for district in data["district"].unique():
    #     district_data = data[data["district"] == district]

    for _, dist_map in districts_map.items():
        ws.merge_cells(f"A{current_row}:G{current_row}")
        ws[f"A{current_row}"] = dist_map["title"]
        ws[f"A{current_row}"].fill = bg_gray
        ws[f"A{current_row}"].font = Font(size=14, bold=True)
        ws[f"A{current_row}"].alignment = align_center

        current_row += 1

        for dist in dist_map["npu_departments"].keys():
            dist_depart_map = dist_map["npu_departments"][dist]

            ws.merge_cells(f"A{current_row}:G{current_row}")
            ws[f"A{current_row}"] = dist_depart_map["title"]
            ws[f"A{current_row}"].fill = bg_gray
            ws[f"A{current_row}"].font = Font(size=12, bold=True)
            ws[f"A{current_row}"].alignment = align_center

            current_row += 1

            start_row_data = current_row
            print(dist_depart_map["departments"])

            for department in dist_depart_map["departments"]:
                district_data = data[data["department"] == department]
                data_fields = [
                    "department",
                    "all_cameras",
                    "all_plates",
                    "all_face",
                    "unavailable",
                    "unavailable_plates",
                    "unavailable_face",
                ]

                for _, row_data in district_data.iterrows():
                    for i in range(1, 8):
                        ws.cell(
                            current_row, i, row_data[data_fields[i - 1]]
                        ).style = "cell_style"

                    for col in range(5, 8):
                        ws.cell(current_row, col).fill = bg_red

                    current_row += 1

        for i in range(start_row, current_row - 1):
            cell_A = ws[f"A{i}"]
            cell_A.border = Border(
                left=thick_line,
                right=cell_A.border.right,
                top=cell_A.border.top,
                bottom=cell_A.border.bottom,
            )
            cell_G = ws[f"G{i}"]
            cell_G.border = Border(
                right=thick_line,
                left=cell_G.border.left,
                top=cell_G.border.top,
                bottom=cell_G.border.bottom,
            )

        start_row = current_row

        for col in range(2, 7):
            ws.cell(row=current_row - 1, column=col).border = Border(
                bottom=thick_line
            )
        ws[f"A{current_row - 1}"].border = Border(
            left=thick_line, bottom=thick_line
        )
        ws[f"G{current_row - 1}"].border = Border(
            right=thick_line, bottom=thick_line
        )

        # ws.cell(current_row, 1, "ВСЬОГО").font = f_bold
        # ws.cell(current_row, 1).fill = bg_gray

        # for col in range(2, 8):
        #     col_letter = get_column_letter(col)
        #     ws.cell(
        #         current_row,
        #         col,
        #         f"SUM({col_letter}{start_row_data}:{col_letter}{current_row - 1})",
        #     )
        #     ws.cell(current_row, col).font = f_bold
        #     ws.cell(current_row, col).fill = bg_gray

        # current_row += 2

    wb.save(out_path)
