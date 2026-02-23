import json
from datetime import datetime
import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    Border,
)
from openpyxl.utils import get_column_letter

from source.camera_info_grabber import CONFIG_DIR
from source.report_styles import openpyxl_styles


def generate_report_openpyxl(data: pd.DataFrame, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = datetime.now().strftime("%d.%m.%Y")
    styles = openpyxl_styles(wb)
    f_title = styles["f_title"]
    font_bold_14 = styles["font_bold_14"]
    font_bold_12 = styles["font_bold_12"]
    align_center = styles["align_center"]

    bg_gray = styles["bg_gray"]
    bg_red = styles["bg_red"]

    thin_line = styles["thin_line"]
    thick_line = styles["thick_line"]
    thin_border = styles["thin_border"]

    ws.column_dimensions["A"].width = 36
    ws.row_dimensions[1].height = 60
    for col in range(2, 8):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 12

    ws.merge_cells("A1:G1")
    ws["A1"] = (
        "Відомості щодо стану функціонування системи відеоспостереження "
        'та відеоаналітики ("Безпечна Чернігівщина")'
    )
    ws["A1"].font = f_title
    ws["A1"].alignment = align_center
    ws["A1"].border = Border(
        left=thick_line, top=thick_line, bottom=thin_line
    )
    ws["G1"].border = Border(
        right=thick_line, top=thick_line, bottom=thin_line
    )
    for col in range(2, 7):
        ws.cell(row=1, column=col).border = Border(
            top=thick_line, bottom=thin_line
        )

    cell_A2 = ws.cell(2, 1, value=datetime.now().strftime("%d.%m.%Y"))
    cell_A2.font = font_bold_14
    cell_A2.border = Border(
        left=thick_line,
        top=thin_border.top,
        right=thin_border.right,
        bottom=thin_border.bottom,
    )

    ws.merge_cells("B2:D2")
    ws["B2"] = "Встановлено"
    ws["B2"].font = font_bold_14
    for i in range(2, 5):
        ws.cell(2, i).border = thin_border

    ws.merge_cells("E2:G2")
    ws["E2"] = "Не працює"
    ws["E2"].font = font_bold_14
    ws["E2"].fill = bg_red
    for i in range(5, 7):
        ws.cell(2, i).border = thin_border

    ws["G2"].border = Border(
        right=thick_line, left=thin_line, top=thin_line, bottom=thin_line
    )

    headers = ["всього", "номерні", "обличчя", "всього", "номерні", "обличчя"]
    ws["A3"] = "Населений пункт"
    ws["A3"].font = font_bold_14
    ws["A3"].border = Border(left=thick_line, bottom=thick_line)
    ws["G3"].border = Border(
        left=thin_line, right=thick_line, bottom=thick_line
    )
    for i, value in enumerate(headers):
        ws.cell(row=3, column=2 + i, value=value).font = font_bold_14
        if i > 2:
            ws.cell(row=3, column=2 + i).fill = bg_red
    for col in range(2, 7):
        ws.cell(row=3, column=col).border = Border(
            left=thin_line, bottom=thick_line
        )

    for row in ws.iter_rows(min_row=1, max_row=3, min_col=1, max_col=7):
        for cell in row:
            cell.alignment = align_center

    current_row = 4
    start_row = 4

    with open(CONFIG_DIR / "region.json", "r", encoding="utf-8") as f:
        districts_map = json.load(f)

    region_stats = []

    for _, dist_map in districts_map.items():
        ws.merge_cells(f"A{current_row}:G{current_row}")
        ws[f"A{current_row}"] = dist_map["title"]
        ws[f"A{current_row}"].fill = bg_gray
        ws[f"A{current_row}"].font = font_bold_14
        ws[f"A{current_row}"].alignment = align_center
        for i in range(1, 8):
            ws.cell(current_row, i).border = Border(
                top=thin_line, bottom=thin_line
            )

        district_deps = []
        current_row += 1

        for dist in dist_map["npu_departments"].keys():
            dist_depart_map = dist_map["npu_departments"][dist]

            ws.merge_cells(f"A{current_row}:G{current_row}")
            ws[f"A{current_row}"] = dist_depart_map["title"]
            ws[f"A{current_row}"].fill = bg_gray
            ws[f"A{current_row}"].font = font_bold_12
            ws[f"A{current_row}"].alignment = align_center
            for i in range(1, 8):
                ws.cell(current_row, i).border = Border(
                    top=thin_line, bottom=thin_line
                )

            current_row += 1
            department_start_row = current_row

            for department in dist_depart_map["departments"]:
                district_data = data[data["department"] == department]
                data_fields = [
                    "department",
                    "all_cameras",
                    "all_plates",
                    "all_face",
                    "unavailable",
                    "unavailable_plates",
                    "unavailable_face",
                ]

                for _, row_data in district_data.iterrows():
                    row_values = [row_data[field] for field in data_fields]
                    ws.append(row_values)
                    for i in range(1, 8):
                        current_cell = ws.cell(current_row, i)
                        current_cell.style = "cell_style"
                        if i > 1:
                            current_cell.alignment = align_center

                    for col in range(5, 8):
                        ws.cell(current_row, col).fill = bg_red

                    current_row += 1

            district_deps.append(current_row)
            cell_A = ws.cell(current_row, 1, "ВСЬОГО")
            cell_A.font = font_bold_12
            cell_A.fill = bg_gray
            cell_A.border = Border(
                left=cell_A.border.left,
                right=thin_line,
                top=thin_line,
                bottom=thin_line,
            )
            for i in range(2, 7):
                current_cell = ws.cell(current_row, i)
                current_cell.border = Border(
                    right=thin_line,
                    left=current_cell.border.left,
                    top=current_cell.border.top,
                    bottom=current_cell.border.bottom,
                )
                if i > 4:
                    current_cell.fill = bg_red
            ws.cell(current_row, 7).fill = bg_red

            for col in range(2, 8):
                current_col = get_column_letter(col)
                ws.cell(
                    current_row,
                    col,
                    value=f"=SUM({current_col}{department_start_row}:{current_col}{current_row - 1})",
                )
                ws.cell(current_row, col).alignment = align_center
                ws.cell(current_row, col).font = font_bold_12
                if col <= 4:
                    ws.cell(current_row, col).fill = bg_gray

            current_row += 1

        region_stats.append(current_row)

        cell_dist = ws.cell(current_row, 1, "ВСЬОГО РАЙОН")
        cell_dist.font = font_bold_12
        cell_dist.fill = bg_gray
        cell_dist.border = Border(
            left=cell_dist.border.left,
            right=thin_line,
            top=thin_line,
            bottom=thin_line,
        )

        for col in range(2, 8):
            current_col = get_column_letter(col)
            formula = [f"{current_col}{dep}" for dep in district_deps]

            ws.cell(current_row, col, value=f"=SUM({','.join(formula)})")
            ws.cell(current_row, col).alignment = align_center
            ws.cell(current_row, col).font = font_bold_12

        for col in range(2, 7):
            ws.cell(current_row, col).border = Border(
                top=thin_line, right=thin_line
            )
        cell_A = ws.cell(current_row, 1)
        cell_A.border = Border(
            left=thick_line,
            top=cell_A.border.top,
            right=cell_A.border.right,
            bottom=cell_A.border.bottom,
        )
        ws.cell(current_row, 7).border = Border(
            top=thin_line, right=thick_line
        )

        for i in range(5, 8):
            current_cell = ws.cell(current_row, i)
            current_cell.fill = bg_red

        for i in range(start_row, current_row):
            cell_A = ws[f"A{i}"]
            cell_A.border = Border(
                left=thick_line,
                right=cell_A.border.right,
                top=cell_A.border.top,
                bottom=cell_A.border.bottom,
            )
            cell_G = ws[f"G{i}"]
            cell_G.border = Border(
                right=thick_line,
                left=cell_G.border.left,
                top=cell_G.border.top,
                bottom=cell_G.border.bottom,
            )

        current_row += 1

        start_row = current_row

    cell_region = ws.cell(current_row, 1, "ВСЬОГО ОБЛАСТЬ")
    cell_region.font = Font(name="Times New Roman", size=12, bold=True)
    cell_region.fill = bg_gray
    for col in range(2, 8):
        current_col = get_column_letter(col)

        region_formula = [f"{current_col}{row}" for row in region_stats]

        ws.cell(current_row, col, value=f"=SUM({','.join(region_formula)})")
        ws.cell(current_row, col).alignment = align_center
        ws.cell(current_row, col).font = font_bold_12
        if col >= 5:
            ws.cell(current_row, col).fill = bg_red

    for col in range(2, 7):
        ws.cell(row=current_row, column=col).border = Border(
            bottom=thick_line, top=thin_line, right=thin_line
        )
    ws[f"A{current_row}"].border = Border(
        left=thick_line, bottom=thick_line, right=thin_line
    )
    ws[f"G{current_row}"].border = Border(
        right=thick_line, bottom=thick_line, top=thin_line
    )

    wb.save(out_path)
